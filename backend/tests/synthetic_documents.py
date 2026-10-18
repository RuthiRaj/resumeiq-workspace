"""Synthetic PDF/DOCX builders for the B2 tests.

Everything here is generated in code and uses clearly fictional content
("Example Candidate", "Example Payments Ltd", example.com addresses). No real
personal data is ever committed to this repository (AGENTS.md, task brief).

The PDF builder writes a minimal but valid PDF by hand: a catalog, a page tree,
Helvetica text pages, a real xref table and a trailer. That keeps the test
fixtures dependency-free (no reportlab) while exercising pypdf's real text
extraction path.
"""

from __future__ import annotations

import io
import zipfile

import docx

TextBlock = str | list[list[str]]


def _escape_pdf_text(text: str) -> str:
    return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def _content_stream(lines: list[str]) -> bytes:
    if not lines:
        return b"q Q"
    body = ["BT", "/F1 12 Tf", "14 TL", "72 740 Td"]
    for line in lines:
        body.append(f"({_escape_pdf_text(line)}) Tj")
        body.append("T*")
    body.append("ET")
    return "\n".join(body).encode("latin-1", "replace")


def build_pdf(pages: list[list[str]]) -> bytes:
    """A text PDF with one content stream per page."""
    page_count = len(pages)
    font_obj = 3 + 2 * page_count
    objects: dict[int, bytes] = {}
    kids = " ".join(f"{3 + 2 * index} 0 R" for index in range(page_count))
    objects[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[2] = f"<< /Type /Pages /Kids [{kids}] /Count {page_count} >>".encode()
    objects[font_obj] = (
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"
    )
    for index, lines in enumerate(pages):
        page_obj = 3 + 2 * index
        content_obj = page_obj + 1
        content = _content_stream(lines)
        objects[page_obj] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_obj} 0 R >> >> /Contents {content_obj} 0 R >>"
        ).encode()
        objects[content_obj] = (
            b"<< /Length "
            + str(len(content)).encode()
            + b" >>\nstream\n"
            + content
            + b"\nendstream"
        )

    out = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for number in sorted(objects):
        offsets[number] = len(out)
        out += f"{number} 0 obj\n".encode() + objects[number] + b"\nendobj\n"

    xref_position = len(out)
    size = max(objects) + 1
    out += f"xref\n0 {size}\n".encode()
    out += b"0000000000 65535 f \n"
    for number in range(1, size):
        if number in offsets:
            out += f"{offsets[number]:010d} 00000 n \n".encode()
        else:
            out += b"0000000000 65535 f \n"
    out += f"trailer\n<< /Size {size} /Root 1 0 R >>\nstartxref\n{xref_position}\n%%EOF\n".encode()
    return bytes(out)


def build_scanned_pdf(pages: int = 1) -> bytes:
    """A PDF with no text layer — what a scanned/image-only upload looks like."""
    return build_pdf([[] for _ in range(pages)])


def build_encrypted_pdf(pages: int = 1) -> bytes:
    """A password-protected PDF: valid, but we never guess the password."""
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(io.BytesIO(build_pdf([["Example Candidate"]] * pages)))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt("example-password")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def build_docx(blocks: list[TextBlock]) -> bytes:
    """A DOCX from paragraphs (``str``) and tables (``list[list[str]]``)."""
    document = docx.Document()
    for block in blocks:
        if isinstance(block, str):
            document.add_paragraph(block)
        elif block:
            table = document.add_table(rows=len(block), cols=max(len(row) for row in block))
            for row_index, row in enumerate(block):
                for cell_index, value in enumerate(row):
                    table.cell(row_index, cell_index).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_zip_with_entries(extra_entries: int, *, as_docx: bool = True) -> bytes:
    """A ZIP archive with a controlled number of entries (zip-bomb fixture)."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        if as_docx:
            archive.writestr("[Content_Types].xml", "<Types/>")
            archive.writestr("word/document.xml", "<w:document/>")
        for index in range(extra_entries):
            archive.writestr(f"word/junk{index}.xml", "x" * 4096)
    return buffer.getvalue()


def build_high_ratio_docx() -> bytes:
    """A DOCX-shaped ZIP whose single payload would expand ~a thousand times."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<w:document/>")
        archive.writestr("word/media/payload.bin", "0" * (4 * 1024 * 1024))
    return buffer.getvalue()
